"""Informe técnico · portada, resumen, introducción, marco conceptual y datos (secciones 1 a 3)."""
from __future__ import annotations

import numpy as np
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from informe_datos import f1, n_es, pval
from informe_lib import HEX_NAVY, NAVY, Informe, _ancho_fijo, _bordes_tabla, _margen_celdas, _runs_con_formato, _sombrear
from enoe import config as C

FIG = C.FIGURES


def portada(inf: Informe, D) -> None:
    d = inf.doc
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("UNIVERSIDAD DE ESPECIALIDADES UNE · PLANTEL CENTRO")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), NAVY
    p = d.add_paragraph()
    r = p.add_run("Ingeniería en Computación · Optativa III: Ciencia de Datos · 9.º semestre")
    r.font.size, r.font.color.rgb = Pt(10.5), RGBColor(0x5F, 0x63, 0x68)
    for _ in range(3):
        d.add_paragraph()
    t = d.add_table(rows=1, cols=1)
    _bordes_tabla(t, interno_h=False)
    _ancho_fijo(t, [16.5])
    _margen_celdas(t, 420, 420, 380, 380)
    c = t.rows[0].cells[0]
    _sombrear(c, HEX_NAVY)
    q = c.paragraphs[0]
    q.paragraph_format.space_after = Pt(10)
    r = q.add_run("Radiografía socioeconómica de México")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(30), RGBColor(255, 255, 255)
    q = c.add_paragraph()
    q.paragraph_format.space_after = Pt(14)
    r = q.add_run("Escolaridad, ingreso y brecha de género")
    r.font.size, r.font.color.rgb = Pt(19), RGBColor(0xCA, 0xDC, 0xFC)
    q = c.add_paragraph()
    q.paragraph_format.space_after = Pt(0)
    r = q.add_run("Informe técnico final del Proyecto Integrador de Ciencia de Datos\nEncuesta Nacional de Ocupación y Empleo (ENOE), 4.º trimestre de 2025 · INEGI")
    r.font.size, r.font.color.rgb = Pt(12), RGBColor(255, 255, 255)
    for _ in range(4):
        d.add_paragraph()
    inf.p("**Equipo 1**", size=13, color=NAVY, despues=4)
    for nombre in ["Juan Evelio Culin Grajeda", "Angel Yireh Garcia Jara", "Luis Andre Gutierrez Zuñiga"]:
        inf.p(nombre, size=13, despues=2)
    for _ in range(3):
        d.add_paragraph()
    inf.p("Documento que explica el proyecto completo: pregunta, datos, limpieza, ingeniería de datos, métodos estadísticos, resultados, "
          "interpretación, limitaciones y guía para reproducirlo.", size=10.5, color=RGBColor(0x5F, 0x63, 0x68), despues=10)
    inf.p("Octubre de 2026", size=11, color=NAVY, negrita=True)


def resumen(inf: Informe, D) -> None:
    R, bn, tot, rt = D.R, D.bn, D.tot, D.rt
    inf.titulo_preliminar("Resumen")
    inf.p(
        f"Este trabajo documenta el Proyecto Integrador de Ciencia de Datos del Equipo 1: un ciclo completo de descubrimiento de conocimiento en bases de datos (KDD) "
        f"aplicado a la Encuesta Nacional de Ocupación y Empleo (ENOE) del INEGI, 4.º trimestre de 2025 (419,038 registros). La pregunta de investigación es cuánto aumenta el "
        f"ingreso laboral por cada año adicional de escolaridad en México, cómo cambia esa relación entre las 32 entidades federativas y si existe una brecha salarial "
        f"significativa entre hombres y mujeres con la misma escolaridad.", alinear="justificado")
    inf.p(
        f"Los datos se limpiaron con reglas documentadas y probadas, se validaron contra siete cifras oficiales del INEGI (las siete coinciden), se almacenaron en un modelo relacional "
        f"(SQLite) y en uno documental (MongoDB) y se analizaron con técnicas que respetan el diseño muestral complejo de la encuesta (bootstrap de Rao-Wu con {R['B']:,} réplicas). "
        f"Se encontró además que el «ingreso cero» del archivo no significa ausencia de ingreso: {D.B_t:.0f} % de los ocupados solo declara un rango de salarios mínimos y {D.D_t:.0f} % no especifica.",
        alinear="justificado")
    inf.p(
        f"**Resultados.** Cada año de escolaridad aumenta el ingreso por hora {f1(rt['nacional_pct']['est'])} % (IC 95 %: {f1(rt['nacional_pct']['ic_inf'])} a {f1(rt['nacional_pct']['ic_sup'])}) y el rendimiento "
        f"es prácticamente igual en estados ricos y pobres (p = {D.dq['p']:.2f}). La brecha de género es de {f1(tot['brecha_hora_pct'])} % por hora y de {f1(tot['brecha_mens_pct'])} % por mes; "
        f"la diferencia mensual se debe sobre todo a que las mujeres trabajan {f1(tot['menos_horas_mujeres_pct'])} % menos horas pagadas. Entre personas comparables la brecha por hora ronda "
        f"{D.g5['brecha_pct']:.0f} % y se concentra en la parte baja de la distribución. La informalidad laboral es de {f1(D.nac['est'])} %, varía de {D.mn_til['til1_est']:.0f} % ({D.mn_til['entidad']}) a "
        f"{D.mx_til['til1_est']:.0f} % ({D.mx_til['entidad']}) y cae de {D.R['til1']['por_nivel']['1']:.0f} % a {D.R['til1']['por_nivel']['4']:.0f} % conforme aumenta la escolaridad.",
        alinear="justificado")
    inf.p("Los resultados describen **asociaciones**, no efectos causales; el documento discute las limitaciones (selección, medición del ingreso, un solo trimestre) y la forma en que se evaluó la robustez de cada conclusión.",
          alinear="justificado")
    inf.p("**Palabras clave:** ENOE; escolaridad; ingreso por hora; brecha de género; informalidad laboral; diseño muestral complejo; bootstrap de Rao-Wu; Oaxaca-Blinder; SQLite; MongoDB.",
          size=10, color=RGBColor(0x5F, 0x63, 0x68), antes=6)


def seccion1(inf: Informe, D) -> None:
    inf.h1("1. Introducción")
    inf.h2("1.1 Contexto y motivación")
    inf.p("La educación suele considerarse la principal vía de movilidad social, pero el mercado laboral mexicano combina una alta informalidad (más de la mitad de las personas ocupadas) con "
          "diferencias persistentes de ingreso entre regiones y entre mujeres y hombres. Saber cuánto «paga» un año adicional de estudio, si ese pago es parejo en todo el país y si mujeres y hombres "
          "con la misma preparación reciben el mismo ingreso son preguntas que orientan decisiones de política educativa, laboral y de igualdad de género.", alinear="justificado")
    inf.p("La **Encuesta Nacional de Ocupación y Empleo (ENOE)** del INEGI es la principal fuente de información del mercado laboral mexicano: es una encuesta continua que cada trimestre entrevista a cientos de miles "
          "de personas y, mediante factores de expansión, permite estimar cifras para todo el país y para cada entidad federativa. Sus microdatos son públicos y anonimizados.", alinear="justificado")
    inf.h2("1.2 Pregunta de investigación")
    inf.recuadro("Pregunta", ["¿Cuánto aumenta el ingreso laboral por cada año adicional de escolaridad en México, cómo cambia esa relación entre entidades federativas y existe una brecha salarial "
                              "significativa entre hombres y mujeres con la misma escolaridad?"])
    inf.h2("1.3 Objetivos")
    inf.p("**Objetivo general.** Cuantificar la relación entre escolaridad e ingreso laboral y la brecha salarial de género en la población ocupada de México, por entidad federativa, a partir de los microdatos de la ENOE.")
    inf.p("**Objetivos específicos.**", despues=2, keep=True)
    inf.numerada(["Estimar la correlación y el retorno de un año adicional de escolaridad sobre el ingreso por hora en cada una de las 32 entidades, comparando estados de mayor y menor ingreso.",
                  "Determinar si existe una diferencia estadísticamente significativa de ingreso por hora entre hombres y mujeres con el mismo nivel de instrucción.",
                  "Identificar las entidades con mayor tasa de informalidad laboral y medir su relación con la escolaridad promedio."])
    inf.h2("1.4 Métricas de éxito")
    inf.p("El protocolo del proyecto fijó tres métricas verificables, cuyo cumplimiento se evalúa en la sección 8.2:", despues=2, keep=True)
    inf.numerada(["**M1.** Brecha salarial de género (%) en el ingreso por hora para cada nivel de instrucción, con intervalo de confianza del 95 % y prueba de hipótesis (significativa si p < 0.05).",
                  "**M2.** Coeficiente de correlación escolaridad-ingreso (Spearman) para las 32 entidades, con intervalo de confianza por bootstrap, y un ranking de estados defendible.",
                  "**M3.** Tablero con al menos cuatro visualizaciones interpretables por una persona no especialista, alimentado por un pipeline que carga datos limpios a una base NoSQL (MongoDB) consultable por entidad."])
    inf.h2("1.5 Alcance")
    inf.p("El estudio es **descriptivo y asociativo**: mide cuánto se asocia la escolaridad con el ingreso y cuánto difieren los ingresos de mujeres y hombres, pero no estima efectos causales (no se puede asegurar que estudiar un "
          "año más *cause* un aumento de 5 %, porque quienes estudian más pueden diferir en habilidades, redes o contexto familiar). Se analiza un solo trimestre (4.º de 2025) y a las personas ocupadas de 15 años y más "
          "(15 a 65 años para el análisis salarial).", alinear="justificado")
    inf.h2("1.6 Cómo leer este documento")
    inf.p("Este trabajo está pensado para que cualquier persona con conocimientos básicos de estadística —y no solo quien lo hizo— entienda **qué se hizo, por qué y qué significa cada resultado**:", despues=2, keep=True)
    inf.vinetas(["Los recuadros **«En palabras sencillas»** explican cada concepto sin tecnicismos.",
                 "La **sección 2** explica los conceptos (la encuesta, el diseño muestral, los métodos estadísticos, las bases de datos); las **secciones 3 a 6** explican qué se hizo; la **7** muestra qué se encontró; la **8** lo interpreta.",
                 ("Los **anexos** incluyen un glosario, la guía para reproducir todo, los resultados de cada entidad y las preguntas frecuentes sobre el método." if inf.web else "Los **anexos** incluyen un glosario, la guía para reproducir todo, preguntas frecuentes para la defensa y un guion de exposición."),
                 "Todas las cifras de este documento se leen automáticamente de los archivos de resultados del proyecto (`results/`), por lo que coinciden con el cuaderno, el tablero y la presentación."]
                + (["**En la versión web**, las figuras marcadas como *interactivas* muestran sus datos al pasar el cursor; «Ver tabla» da los valores y «Imagen» la versión estática. "
                    "Al elegir una entidad en la barra superior (o con un clic en una figura) se resalta en todas las figuras y en las tablas; los términos con subrayado punteado muestran su definición; la tecla **/** abre el buscador."]
                   if inf.web else []))


def seccion2(inf: Informe, D) -> None:
    R, bn = D.R, D.bn
    inf.h1("2. Marco conceptual: lo que hay que entender")
    inf.p("Antes de ver qué se hizo conviene entender las ideas que sostienen el análisis. Esta sección las explica en lenguaje sencillo, de lo general (el método de trabajo) a lo particular (cada técnica).", alinear="justificado")

    inf.h2("2.1 El ciclo KDD")
    inf.p("**KDD** (*Knowledge Discovery in Databases*) es el proceso completo que convierte datos crudos en conocimiento útil. No es solo «correr un modelo»: la mayor parte del trabajo está antes (elegir y limpiar los datos) "
          "y después (interpretar y comunicar). Tiene cinco fases, que corresponden a las semanas del curso:", alinear="justificado")
    inf.figura(FIG / "D01_kdd.png", 16.2, "Las cinco fases del ciclo KDD y qué se hizo en cada una.")
    inf.tabla([["1. Selección", "1", "Elegir la fuente y el subconjunto de datos relevantes", "Sección 3"],
               ["2. Preprocesamiento", "4", "Limpiar: faltantes, duplicados, tipos, valores atípicos e imposibles", "Sección 4"],
               ["3. Transformación", "4 · 6 · 7", "Crear variables, unir tablas y almacenarlas (relacional y documental)", "Sección 5"],
               ["4. Minería / Análisis", "2 · 3 · 5", "Estadística, pruebas de hipótesis, regresión, visualización", "Secciones 6 y 7"],
               ["5. Interpretación", "5", "Extraer conocimiento, comunicarlo y recomendar", "Secciones 8 y 9"]],
              ["Fase", "Semana", "Qué se hace", "Dónde está"], [4.0, 2.0, 7.7, 2.8], pie="Fases del ciclo KDD y su correspondencia con el documento.", alinear=["izq", "centro", "izq", "izq"])

    inf.h2("2.2 La ENOE y su diseño muestral")
    inf.h3("¿Qué es la ENOE?")
    inf.p("Es una encuesta del INEGI que se levanta de manera continua en las viviendas del país. Cada trimestre el INEGI publica una base de **microdatos**: archivos en los que cada fila es una persona (tabla SDEMT, "
          "sociodemográfica y laboral) o un cuestionario (tablas COE1T y COE2T, ocupación y empleo). En el 4.º trimestre de 2025 la tabla SDEMT trae **419,038 personas** y 115 columnas. "
          "Cada vivienda se entrevista hasta cinco veces consecutivas (panel rotativo, variable `n_ent`).", alinear="justificado")
    inf.h3("Diseño probabilístico, estratificado y por conglomerados")
    inf.p("No se entrevista a una muestra «al azar simple» de personas: sería carísimo recorrer el país. En su lugar se usa un diseño **bietápico, estratificado y por conglomerados**: primero se agrupan áreas "
          "parecidas en **estratos**, luego se sortean **Unidades Primarias de Muestreo (UPM)**, que son conjuntos de viviendas, y dentro de ellas se eligen viviendas y personas. En este trimestre hay 724 estratos y 21,194 UPM.",
          alinear="justificado")
    inf.figura(FIG / "D04_diseno_muestral.png", 16.0, "Niveles del diseño muestral de la ENOE y significado del factor de expansión (fac_tri).")
    inf.h3("El factor de expansión: por qué cada persona pesa distinto")
    inf.p("Cada persona de la base trae un número, `fac_tri`, que indica **a cuántas personas de la población representa**. Las personas de algunas zonas se entrevistan con mayor o menor probabilidad, y el factor corrige eso: "
          "la suma de `fac_tri` en toda la base es 131.0 millones de personas (la población del país) y la de las personas ocupadas de 15 años y más es 59.8 millones. Por eso **todos los promedios, porcentajes y totales de este "
          "trabajo están ponderados**: un promedio sin ponderar describiría a la muestra, no al país.", alinear="justificado")
    inf.h3("Efecto de diseño: la muestra «vale» menos de lo que parece")
    inf.p(f"Las personas de una misma UPM se parecen entre sí (viven en la misma zona, tienen empleos similares), así que aportan menos información que personas elegidas al azar en todo el país. "
          f"El **efecto de diseño** (*deff*) mide cuánto más grande es la varianza real que la de una muestra aleatoria simple: para la tasa de informalidad es de {D.deff.loc['TIL1 nacional', 'deff']:.1f}, "
          f"de modo que los {D.n_u:,} ocupados de la base equivalen a unas {D.deff.loc['TIL1 nacional', 'n_efectivo'] / 1000:.0f} mil observaciones independientes. Ignorarlo haría que los intervalos de confianza fueran "
          f"aproximadamente la mitad de anchos de lo que deberían y que todo pareciera «significativo».", alinear="justificado")
    inf.recuadro("En palabras sencillas", ["Imagina que cada persona encuestada es un «representante» de cientos de personas que no fueron entrevistadas. Por eso hay que contar a cada representante tantas veces como personas representa (ponderar). "
                                           "Y como los representantes de una misma zona se parecen, tener 190 mil entrevistas no equivale a 190 mil opiniones totalmente independientes: equivale a unas 40 mil."])

    inf.h2("2.3 Las variables laborales")
    inf.tabla([["Persona ocupada", "Persona de 15+ años que trabajó al menos una hora la semana de referencia (`clase2 = 1`)."],
               ["Ingreso mensual (`ingocup`)", "Ingreso mensual de la ocupación principal, en pesos. Vale 0 si la persona no dio un monto exacto (ver sección 4.3)."],
               ["Horas trabajadas (`hrsocup`)", "Horas trabajadas la semana pasada."],
               ["Ingreso por hora", "`ingocup / (hrsocup × 4.3)`; 4.3 es el número de semanas por mes que usa el INEGI. Coincide con la variable `ing_x_hrs` del INEGI (error relativo < 0.01 %)."],
               ["Salario mínimo (`salario`)", "Salario mínimo mensual de la zona de la persona: $8,364 (zona general) o $12,596 (frontera norte)."],
               ["Escolaridad (`anios_esc`, `niv_ins`)", "Años de escolaridad aprobados y nivel de instrucción (primaria incompleta, primaria completa, secundaria completa, medio superior y superior)."],
               ["Informalidad (`emp_ppal`)", "Empleo informal en la ocupación principal. **TIL1** = personas ocupadas en empleo informal ÷ personas ocupadas de 15+ años, definición oficial del INEGI."],
               ["Entidad (`cve_ent`)", "Clave de la entidad federativa, 1 a 32."]],
              ["Concepto", "Definición"], [4.8, 11.7], pie="Variables laborales y sus definiciones.", alinear=["izq", "izq"], negrita_col0=True)

    inf.h2("2.4 Estadística en palabras sencillas")
    inf.p("Cada técnica se explica con la misma estructura: *la idea*, *cómo se calcula* y *cómo se lee*.", despues=4)

    inf.h3("Logaritmo natural y media geométrica")
    inf.p("**La idea.** El ingreso por hora tiene una «cola larga»: la mayoría gana poco y unas pocas personas ganan muchísimo, lo que distorsiona los promedios. Con el logaritmo natural, esa cola se acorta y la distribución se "
          f"vuelve casi simétrica (la asimetría pasa de {D.dg['asimetria_ingreso_nivel']:.1f} a {D.dg['asimetria_ln_ingreso']:.2f}).", alinear="justificado")
    inf.p("**Cómo se calcula.** Se trabaja con ln(ingreso por hora). Una propiedad útil: la diferencia entre dos logaritmos es, aproximadamente, una diferencia porcentual (ln 110 − ln 100 = 0.095 ≈ 10 %). "
          "El promedio de los logaritmos, devuelto a pesos con la exponencial, es la **media geométrica**.", alinear="justificado")
    inf.p("**Cómo se lee.** Cuando se dice «las mujeres ganan 3.7 % menos por hora» se comparan las medias geométricas de ingreso por hora de mujeres y hombres.", alinear="justificado")

    inf.h3("Correlación de Spearman")
    inf.p("**La idea.** Mide si dos variables «suben juntas»: si las personas con más escolaridad tienden a tener más ingreso, sin importar la forma exacta de la relación ni los valores extremos.", alinear="justificado")
    inf.p("**Cómo se calcula.** Se reemplaza cada valor por su *lugar* (rango) entre todos los valores (el más bajo es 1) y se calcula la correlación de Pearson entre los rangos. Los empates reciben el rango medio. "
          "En este proyecto cada persona cuenta tantas veces como su `fac_tri` (rangos medios ponderados). El ejemplo siguiente usa seis personas ilustrativas.", alinear="justificado")
    inf.figura(FIG / "D08_spearman_ejemplo.png", 15.5, "Ejemplo ilustrativo (seis personas): al convertir los valores en rangos, la correlación de Spearman resulta 0.93.")
    inf.p("**Cómo se lee.** Va de −1 (relación inversa perfecta) a +1 (perfecta), y 0 significa que no hay relación monótona. Como referencia común, alrededor de 0.1 es débil, 0.3 moderada y 0.5 fuerte. "
          f"La correlación nacional entre escolaridad e ingreso por hora es {D.rt['spearman_nacional']['est']:.2f} (moderada).", alinear="justificado")

    inf.h3("Regresión lineal y ecuación de Mincer")
    inf.p("**La idea.** Una regresión estima cuánto cambia una variable (aquí, el ingreso) cuando cambia otra (los años de escolaridad), *manteniendo constantes* las demás (edad, sexo, entidad...). "
          "La **ecuación de Mincer** (1974) es el modelo clásico de la economía laboral para medir el rendimiento de la educación.", alinear="justificado")
    inf.formula("ln(ingreso por hora) = α + β · años de escolaridad + γ₁ · edad + γ₂ · edad² + δ · mujer + efectos de entidad y de localidad + error")
    inf.p("**Cómo se lee.** El coeficiente β mide cuánto aumenta el *logaritmo* del ingreso por cada año de escolaridad; el aumento porcentual es 100 × (e^β − 1). Si β = 0.0501, un año más de estudio se asocia con "
          f"+{f1(D.rt['nacional_pct']['est'])} % de ingreso por hora; dos años, con (1.051)² − 1 = +10.5 %. La edad entra al cuadrado porque el ingreso sube con la experiencia y luego se aplana.", alinear="justificado")

    inf.h3("Brecha salarial de género")
    inf.p("**La idea y el cálculo.** Se compara el ingreso medio geométrico de mujeres y hombres:", despues=2, keep=True)
    inf.formula("brecha (%) = 100 × [ 1 − (ingreso medio geométrico de las mujeres ÷ ingreso medio geométrico de los hombres) ]")
    inf.p(f"**Ejemplo con los datos.** Los hombres ganan en promedio geométrico ${D.tot['ing_hora_h']:.2f} por hora y las mujeres ${D.tot['ing_hora_m']:.2f}: la brecha es 1 − {D.tot['ing_hora_m']:.2f}/{D.tot['ing_hora_h']:.2f} = {f1(D.tot['brecha_hora_pct'])} %. "
          "Un valor negativo indicaría que las mujeres ganan más.", alinear="justificado")

    inf.h3("Pruebas de hipótesis, valor p e intervalo de confianza")
    inf.p("**La idea.** Como trabajamos con una muestra, toda diferencia observada podría deberse al azar. Una **prueba de hipótesis** plantea una *hipótesis nula* (H₀: no hay brecha) y calcula el **valor p**: "
          "la probabilidad de observar una diferencia tan grande como la que vimos *si en realidad no hubiera brecha*. Si p < 0.05 (nivel α = 5 %) se rechaza H₀ y se dice que la diferencia es **estadísticamente significativa**.", alinear="justificado")
    inf.p("El **intervalo de confianza (IC) del 95 %** da un rango de valores plausibles para la cantidad real: si el estudio se repitiera muchas veces, el 95 % de los intervalos construidos así contendría el valor verdadero. "
          "Si el intervalo **no incluye el 0**, la diferencia es significativa; si lo incluye, no se puede afirmar que exista.", alinear="justificado")
    inf.figura(FIG / "D06_como_leer_ic.png", 15.8, "Cómo leer un intervalo de confianza con tres casos reales: uno que excluye el 0 (hay brecha), uno que lo toca (evidencia débil) y uno que lo incluye (no hay brecha demostrable).")

    inf.h3("Comparaciones múltiples: la corrección de Holm")
    inf.p("Si se hacen muchas pruebas, algunas saldrán «significativas» solo por azar. La corrección de **Holm** (1979) controla la probabilidad de equivocarse al menos una vez en todo el conjunto: "
          "ordena los valores p de menor a mayor y los compara con umbrales cada vez menos exigentes. Se aplica a las cuatro pruebas por nivel de instrucción y a las categorías de cada dimensión de heterogeneidad.", alinear="justificado")

    inf.h3("Tamaño del efecto: la d de Cohen")
    inf.p("Con muchas observaciones casi cualquier diferencia resulta «significativa». Para saber si importa en la práctica se mide su tamaño: la **d de Cohen** es la diferencia de medias dividida entre la desviación estándar "
          "(0.2 se considera pequeña, 0.5 mediana y 0.8 grande). Significativo no significa grande.", alinear="justificado")

    inf.h3("Bootstrap de Rao-Wu y linealización de Taylor")
    inf.p("**La idea.** Para medir la incertidumbre (el error estándar) hay que saber cuánto cambiaría el resultado si se repitiera la encuesta. No se puede repetir, pero se puede *simular*: el **bootstrap** vuelve a muestrear "
          "los propios datos. Como la ENOE es por conglomerados, lo que se remuestrea no son personas sino **UPM completas**, respetando los estratos. El método de **Rao y Wu (1988)** lo hace de forma que reproduce correctamente la varianza del diseño.", alinear="justificado")
    inf.figura(FIG / "D05_rao_wu.png", 16.2, "El bootstrap de Rao-Wu en cuatro pasos (ejemplo con un estrato de cuatro UPM).")
    inf.p("**Cómo se usa.** Con 1,000 réplicas, la desviación estándar de las estimaciones replicadas es el error estándar y los percentiles 2.5 y 97.5 son el intervalo de confianza del 95 %. "
          "Para los modelos de regresión y logit se usa en cambio la **linealización de Taylor**, una fórmula analítica equivalente que también respeta estratos y UPM. En las pruebas del proyecto, ambos métodos coinciden (dentro de 8 %).", alinear="justificado")

    inf.h3("Descomposición de Oaxaca-Blinder")
    inf.p("**La idea.** Separa la brecha entre dos grupos en dos partes: la **explicada** (los grupos difieren en sus características: escolaridad, edad, lugar de residencia, tipo de empleo...) y la **no explicada** "
          "(a las *mismas* características se les paga distinto). Se usa la referencia agrupada de Neumark (1988): los coeficientes de comparación salen de una regresión con hombres y mujeres juntos.", alinear="justificado")
    inf.figura(FIG / "D07_oaxaca_idea.png", 15.8, "Idea de la descomposición de Oaxaca-Blinder con los resultados del proyecto (puntos logarítmicos, hombres − mujeres).")
    inf.p("**Cómo se lee.** La parte no explicada **no prueba discriminación**: incluye todo lo que no se pudo medir (ocupación específica, experiencia real, negociación, selección).", alinear="justificado")

    inf.h3("No respuesta: imputación por rangos y ponderación por probabilidad inversa (IPW)")
    inf.p("Cuando una parte de la muestra no responde y quienes no responden son distintos de quienes sí, los resultados se sesgan. Se usan dos correcciones: **imputar** el ingreso de quien solo dio un rango (asignándole "
          "el ingreso típico de quienes ganan en ese rango) y **IPW**, que da más peso a quienes se parecen a los que no respondieron (peso = 1 ÷ probabilidad de responder, estimada con una regresión logística).", alinear="justificado")

    inf.h3("Regresión logística")
    inf.p("Cuando la variable que se explica es sí/no (por ejemplo, tener un empleo informal), se usa un **logit**. Sus coeficientes se leen como **razones de momios** (*odds ratio*): un valor de 0.82 por año de escolaridad "
          "significa que cada año reduce 18 % los momios de ser informal. Para una lectura más directa se reporta el **efecto marginal promedio**: cuántos puntos porcentuales cambia la probabilidad.", alinear="justificado")

    inf.h2("2.5 Ingeniería de datos")
    inf.vinetas(["**ETL** (*Extract, Transform, Load*): proceso que extrae los datos de su fuente, los transforma (limpia, une, calcula variables) y los carga en un almacén. Aquí es un programa en Python que se puede ejecutar de principio a fin.",
                 "**Base de datos relacional (SQL):** guarda los datos en tablas con columnas de tipos fijos y relaciones entre tablas. Es ideal para datos tabulares y para agregaciones (`GROUP BY`, `JOIN`). Las restricciones (`PRIMARY KEY`, `FOREIGN KEY`, `CHECK`) impiden guardar datos inconsistentes.",
                 "**Esquema en estrella:** una tabla central de hechos (aquí, una fila por persona ocupada) rodeada de tablas de dimensiones pequeñas (entidad, sexo, nivel de instrucción...).",
                 "**Base de datos documental (NoSQL):** guarda documentos JSON anidados, sin esquema rígido. Es ideal para servir «la ficha completa» de un elemento (una entidad) con una sola lectura.",
                 "**Persistencia políglota:** usar cada tipo de base donde mejor funciona. Aquí, SQLite para analizar y MongoDB para servir la ficha de cada entidad.",
                 "**Idempotencia y determinismo:** ejecutar el proceso dos veces produce exactamente el mismo resultado (se fija una semilla para todo lo aleatorio y se recrea la base desde cero)."])


def seccion3(inf: Informe, D) -> None:
    inf.h1("3. Datos (KDD fase 1: Selección)")
    inf.h2("3.1 Fuente y archivos")
    inf.p("Se usan los microdatos de la ENOE del 4.º trimestre de 2025, publicados por el INEGI (https://www.inegi.org.mx/programas/enoe/15ymas/) en archivos CSV con codificación latin-1. "
          "El paquete incluye cinco tablas; el proyecto usa dos y guarda la huella SHA-256 de cada archivo para poder demostrar con qué datos se corrió todo.", alinear="justificado")
    filas = [[a["tabla"], a["archivo"], f"{a['filas']:,}", str(a["columnas_leidas"]), f"{a['bytes'] / 1e6:.0f}", a["sha256"][:16] + "…"] for a in D.meta["archivos_origen"]]
    inf.tabla(filas, ["Tabla", "Archivo", "Filas", "Columnas leídas", "MB", "SHA-256 (inicio)"], [1.8, 4.2, 2.2, 2.6, 1.3, 4.4],
              pie="Tablas de origen utilizadas. SDEMT trae 115 columnas (se leen las 35 necesarias).", alinear=["izq", "izq", "der", "der", "der", "izq"])
    inf.vinetas(["**SDEMT** (sociodemográfica y laboral): una fila por residente del hogar, de todas las edades; trae el ponderador, el diseño muestral, la escolaridad, el ingreso y las horas.",
                 "**COE2T** (cuestionario de ocupación, parte 2): el detalle de cómo se captó el ingreso (monto, o rango de salarios mínimos cuando no hubo monto).",
                 "Las tablas COE1T, HOGT y VIVT no aportan variables al problema y no se usan."])
    inf.h2("3.2 Llave de persona y unión de tablas")
    inf.p("Para unir SDEMT con COE2T hace falta una llave que identifique a cada persona. La llave que se cita habitualmente (`cd_a, ent, con, v_sel, n_hog, h_mud, n_ren`) **no es única** en este trimestre: tiene 8,371 duplicados. "
          "Se comprobó empíricamente que la combinación `cd_a, cve_ent, con, upm, d_sem, n_pro_viv, v_sel, n_hog, h_mud, n_ent, n_ren` sí lo es (0 duplicados en ambas tablas). La unión se hace con la condición «uno a uno»: "
          "si alguna persona ocupada de SDEMT no encontrara pareja en COE2T, el proceso se detiene. Resultado: los 190,187 ocupados encuentran su pareja.", alinear="justificado")
    inf.h2("3.3 Variables utilizadas")
    inf.tabla([["r_def, c_res", "SDEMT", "Resultado de la entrevista; condición de residencia", "r_def = 0 (completa); c_res ∈ {1, 3} (residente habitual)"],
               ["cve_ent", "SDEMT", "Entidad federativa", "1 a 32"],
               ["upm, est_d_tri", "SDEMT", "Unidad primaria de muestreo y estrato de diseño", "21,194 UPM en 724 estratos"],
               ["fac_tri", "SDEMT", "Factor de expansión trimestral", "> 0; suma 131.0 millones"],
               ["sex, eda", "SDEMT", "Sexo y edad", "sex: 1 hombre, 2 mujer; eda: 0 a 98"],
               ["niv_ins, anios_esc", "SDEMT", "Nivel de instrucción y años de escolaridad", "anios_esc = 99 significa «no especificado»"],
               ["clase1, clase2", "SDEMT", "Condición de actividad", "clase2 = 1: ocupado"],
               ["pos_ocu", "SDEMT", "Posición en la ocupación", "1 subordinado, 2 empleador, 3 cuenta propia, 4 sin pago"],
               ["emp_ppal", "SDEMT", "Empleo informal en la ocupación principal", "1 informal, 2 formal"],
               ["rama_est2", "SDEMT", "Gran rama de actividad", "11 ramas"],
               ["t_loc_tri, salario", "SDEMT", "Tamaño de localidad; salario mínimo mensual", "4 tamaños; $8,364 / $12,596"],
               ["hrsocup", "SDEMT", "Horas trabajadas la semana pasada", "0 a 168"],
               ["ingocup, ing_x_hrs, ing7c", "SDEMT", "Ingreso mensual, ingreso por hora y rango en salarios mínimos", "ingocup = 0 si no hay monto exacto"],
               ["p6b2, p6c", "COE2T", "Monto declarado; rango de salarios mínimos", "p6c: 1 a 7 rangos, 8 no sabe, 9 no respondió"]],
              ["Variable", "Tabla", "Significado", "Valores / notas"], [3.2, 1.5, 6.2, 5.6], pie="Variables de origen utilizadas (diccionario completo en docs/diccionario_de_datos.md).",
              alinear=["izq", "izq", "izq", "izq"], fuente=9)
    inf.h2("3.4 Validación externa: ¿las cifras coinciden con las oficiales?")
    inf.p("Antes de analizar nada se comprobó que los filtros y los ponderadores reproducen lo que el INEGI publicó para este trimestre. Si coinciden, se interpretan bien los códigos y los factores de expansión "
          "y los resultados son comparables con la estadística oficial. La base oficial es: personas ocupadas de 15 años y más, entrevista completa, residentes habituales, ponderadas con `fac_tri`.", alinear="justificado")
    filas = [[r.indicador, f"{r.estimado_proyecto:,.2f}", f"{r.oficial_inegi:,.1f}", "Sí" if r.coincide else "No"] for r in D.etl_validacion_externa.itertuples()]
    inf.tabla(filas, ["Indicador", "Proyecto", "INEGI (publicado)", "¿Coincide?"], [7.0, 3.0, 3.6, 2.9], pie="Validación externa: cifras del proyecto frente a las publicadas por el INEGI para el 4T-2025.",
              alinear=["izq", "der", "der", "centro"])
    ef = D.meta["efecto_base_menores_15"]
    inf.p(f"**Las siete comparaciones coinciden** al redondear a la precisión publicada. Un hallazgo de la validación: si se incluyeran los {ef['menores_de_15_ocupados_en_muestra']:,} menores de 15 años que aparecen como ocupados en la muestra, "
          f"el TIL1 subiría de {ef['til1_base_15_y_mas']:.2f} % a {ef['til1_con_menores_de_15']:.2f} % y dejaría de coincidir con el oficial; por eso toda la población de estudio se restringe a 15 años y más.", alinear="justificado")
    inf.recuadro("Nota de transparencia", ["Las cifras oficiales de referencia se obtuvieron de los boletines del INEGI para el 4T-2025 mediante una búsqueda web; conviene contrastarlas con el PDF oficial antes de citarlas."], fondo="FFF7E6", borde="EDA100")
