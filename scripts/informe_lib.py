"""Ayudantes para construir el informe técnico en Word con python-docx (estilos, tablas, figuras, recuadros, campos)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image

NAVY, BLUE, INK, GRIS = RGBColor(0x1E, 0x2A, 0x6B), RGBColor(0x2A, 0x78, 0xD6), RGBColor(0x1B, 0x1B, 0x1F), RGBColor(0x5F, 0x63, 0x68)
HEX_NAVY, HEX_TINTE, HEX_LINEA, HEX_ZEBRA = "1E2A6B", "F3F5FB", "C9D3EE", "F7F8FC"
FUENTE = "Calibri"
TAM_BASE = 10.5
ANCHO_TEXTO_CM = 16.5


# --------------------------------------------------------------------------- #
# utilidades XML
# --------------------------------------------------------------------------- #
def _sombrear(celda, hex_fill: str) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    for viejo in tcPr.findall(qn("w:shd")):
        tcPr.remove(viejo)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _bordes_tabla(tabla, color: str = HEX_LINEA, interno_h: bool = True, exterior: bool = False) -> None:
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(viejo)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        activo = (lado in ("insideH",) and interno_h) or (exterior and lado in ("top", "bottom"))
        e.set(qn("w:val"), "single" if activo else "nil")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _margen_celdas(tabla, arriba=50, abajo=50, izq=90, der=90) -> None:
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(viejo)
    m = OxmlElement("w:tblCellMar")
    for lado, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _ancho_fijo(tabla, anchos_cm: list[float]) -> None:
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(viejo)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    total = int(sum(anchos_cm) / 2.54 * 1440)
    for viejo in tblPr.findall(qn("w:tblW")):
        tblPr.remove(viejo)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(total))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)


def _campo(parrafo, instruccion: str, texto_previo: str = "1") -> None:
    """Inserta un campo de Word (PAGE, NUMPAGES, TOC...)."""
    def run_fld(tipo):
        r = parrafo.add_run()
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), tipo)
        r._r.append(f)
        return r

    run_fld("begin")
    r = parrafo.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    r._r.append(it)
    run_fld("separate")
    parrafo.add_run(texto_previo)
    run_fld("end")


def _runs_con_formato(parrafo, texto: str, size: float | None = None, color: RGBColor | None = None, negrita=False, cursiva=False) -> None:
    """Mini-marcado: **negrita**, *cursiva* y `código`."""
    partes = re.split(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)", texto)
    for p in partes:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            _runs_con_formato(parrafo, p[2:-2], size, color, True, cursiva)      # negrita con código/cursiva anidados
            continue
        b, i, mono = negrita, cursiva, False
        if p.startswith("`") and p.endswith("`") and len(p) > 2:
            p, mono = p[1:-1], True
        elif p.startswith("*") and p.endswith("*") and len(p) > 2:
            p, i = p[1:-1], True
        r = parrafo.add_run(p)
        r.bold, r.italic = b, i
        if mono:
            r.font.name = "Consolas"
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or TAM_BASE) - 1)
        elif size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


# --------------------------------------------------------------------------- #
class Informe:
    def __init__(self) -> None:
        self.doc = Document()
        self.n_fig = 0
        self.n_tab = 0
        self._configurar()

    # ----- configuración -----
    def _configurar(self) -> None:
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)          # Carta
        s.left_margin = s.right_margin = Cm(2.54)
        s.top_margin, s.bottom_margin = Cm(2.3), Cm(2.2)
        s.different_first_page_header_footer = True
        st = d.styles
        n = st["Normal"]
        n.font.name, n.font.size = FUENTE, Pt(TAM_BASE)
        n.font.color.rgb = INK
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.12
        for nombre, tam, color, antes, despues, salto in [("Heading 1", 21, NAVY, 0, 10, True), ("Heading 2", 15, NAVY, 16, 6, False), ("Heading 3", 12.5, BLUE, 12, 4, False)]:
            h = st[nombre]
            h.font.name, h.font.size, h.font.bold, h.font.italic = FUENTE, Pt(tam), True, False
            h.font.color.rgb = color
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
            h.element.rPr.rFonts.set(qn("w:ascii"), FUENTE)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), FUENTE)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(antes), Pt(despues)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.page_break_before = salto
        for nivel, (sangria, negrita, antes) in {1: (0.0, True, 3), 2: (0.6, False, 0)}.items():     # estilos del índice (Word los reutiliza al actualizar el campo)
            e = st.add_style(f"toc {nivel}", WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            e.base_style = st["Normal"]
            e.font.size, e.font.bold = Pt(10), negrita
            e.paragraph_format.space_before, e.paragraph_format.space_after = Pt(antes), Pt(1.5)
            e.paragraph_format.left_indent = Cm(sangria)
            e.paragraph_format.line_spacing = 1.0
        cap = st["Caption"]
        cap.font.name, cap.font.size, cap.font.italic, cap.font.bold = FUENTE, Pt(9.5), False, False
        cap.font.color.rgb = GRIS
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.space_before = Pt(2)
        for nombre in ("List Bullet", "List Bullet 2"):
            lb = st[nombre]
            lb.font.name, lb.font.size = FUENTE, Pt(TAM_BASE)
            lb.paragraph_format.space_after = Pt(3)
        d.core_properties.title = "Radiografía socioeconómica de México: escolaridad, ingreso y brecha de género"
        d.core_properties.author = "Equipo 1: Juan Evelio Culin Grajeda, Angel Yireh Garcia Jara, Luis Andre Gutierrez Zuñia"
        d.core_properties.subject = "Informe técnico final · Proyecto Integrador de Ciencia de Datos · ENOE 2025-T4"
        d.core_properties.language = "es-MX"
        self._pie_pagina(s)

    def _pie_pagina(self, seccion) -> None:
        f = seccion.footer
        p = f.paragraphs[0]
        p.text = ""
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Cm(ANCHO_TEXTO_CM), alignment=2)  # derecha
        r = p.add_run("Informe técnico · Proyecto Integrador de Ciencia de Datos · ENOE 2025-T4\t")
        r.font.size, r.font.color.rgb = Pt(8.5), GRIS
        r2 = p.add_run("Página ")
        r2.font.size, r2.font.color.rgb = Pt(8.5), GRIS
        _campo(p, "PAGE")
        for run in p.runs[2:]:
            run.font.size, run.font.color.rgb = Pt(8.5), GRIS
        # separador superior del pie
        pPr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "6")
        top.set(qn("w:color"), HEX_LINEA)
        bd.append(top)
        pPr.append(bd)

    # ----- bloques de texto -----
    def h1(self, t, salto: bool = True):
        h = self.doc.add_heading(t, level=1)
        if not salto:
            h.paragraph_format.page_break_before = False
            h.paragraph_format.space_before = Pt(22)
        return h

    def h2(self, t):
        return self.doc.add_heading(t, level=2)

    def h3(self, t):
        return self.doc.add_heading(t, level=3)

    def titulo_preliminar(self, t: str, salto: bool = True):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = salto
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(t)
        r.font.size, r.font.bold, r.font.color.rgb = Pt(21), True, NAVY
        return p

    def p(self, texto: str, alinear=None, size=None, color=None, negrita=False, cursiva=False, despues=None, antes=None, keep=False):
        par = self.doc.add_paragraph()
        _runs_con_formato(par, texto, size, color, negrita, cursiva)
        if alinear == "centro":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif alinear == "justificado":
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if despues is not None:
            par.paragraph_format.space_after = Pt(despues)
        if antes is not None:
            par.paragraph_format.space_before = Pt(antes)
        par.paragraph_format.keep_with_next = keep
        return par

    def vinetas(self, items: list[str], nivel: int = 1) -> None:
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet" if nivel == 1 else "List Bullet 2")
            _runs_con_formato(par, it)

    def numerada(self, items: list[str], inicio: int = 1) -> None:
        for k, it in enumerate(items, start=inicio):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.9)
            par.paragraph_format.space_after = Pt(4)
            r = par.add_run(f"{k}.\t")
            r.bold = True
            par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            _runs_con_formato(par, it)

    def formula(self, texto: str) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        _bordes_tabla(t, interno_h=False)
        _ancho_fijo(t, [ANCHO_TEXTO_CM])
        _margen_celdas(t, 90, 90, 160, 160)
        c = t.rows[0].cells[0]
        _sombrear(c, HEX_TINTE)
        par = c.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        _runs_con_formato(par, texto, size=11.5, color=NAVY)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def recuadro(self, titulo: str, lineas: list[str], fondo: str = HEX_TINTE, borde: str = HEX_NAVY) -> None:
        """Recuadro destacado con borde izquierdo (para «En palabras sencillas», advertencias, etc.)."""
        t = self.doc.add_table(rows=1, cols=1)
        _bordes_tabla(t, interno_h=False)
        _ancho_fijo(t, [ANCHO_TEXTO_CM])
        _margen_celdas(t, 100, 100, 200, 160)
        c = t.rows[0].cells[0]
        _sombrear(c, fondo)
        tcPr = c._tc.get_or_add_tcPr()
        bd = OxmlElement("w:tcBorders")
        izq = OxmlElement("w:left")
        izq.set(qn("w:val"), "single")
        izq.set(qn("w:sz"), "30")
        izq.set(qn("w:space"), "0")
        izq.set(qn("w:color"), borde)
        bd.append(izq)
        tcPr.append(bd)
        par = c.paragraphs[0]
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(titulo)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), NAVY
        for k, ln in enumerate(lineas):
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(3 if k < len(lineas) - 1 else 0)
            _runs_con_formato(q, ln, size=10.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def codigo(self, texto: str, size: float = 8.5) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        _bordes_tabla(t, interno_h=False)
        _ancho_fijo(t, [ANCHO_TEXTO_CM])
        _margen_celdas(t, 80, 80, 140, 140)
        c = t.rows[0].cells[0]
        _sombrear(c, "F5F5F7")
        primero = True
        for ln in texto.strip("\n").split("\n"):
            par = c.paragraphs[0] if primero else c.add_paragraph()
            primero = False
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(ln if ln else " ")
            r.font.name, r.font.size = "Consolas", Pt(size)
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.color.rgb = RGBColor(0x2B, 0x2F, 0x3A)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ----- bloques con versión propia en la tesina web (informe_html.InformeHTML): referencias, glosario y preguntas frecuentes -----
    web = False

    def referencia(self, texto: str) -> None:
        """Una referencia bibliográfica con sangría francesa."""
        par = self.p(texto, size=10, despues=4)
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)

    def glosario(self, items: list[tuple[str, str]]) -> None:
        self.tabla([[t, d] for t, d in items], ["Término", "Significado"], [4.6, 11.9], pie="Glosario de términos.", alinear=["izq", "izq"], fuente=9, negrita_col0=True)

    def pregunta(self, q: str, a: str) -> None:
        """Una pregunta frecuente y su respuesta."""
        self.p(f"**{q}**", despues=1, keep=True)
        self.p(a, alinear="justificado", despues=6)

    # ----- tablas y figuras -----
    def tabla(self, filas: list[list], encabezados: list[str], anchos_cm: list[float], pie: str | None = None, fuente: float = 9.5,
              alinear: list[str] | None = None, negrita_col0: bool = False, resaltar_fila: set[int] | None = None, mantener: bool | None = None) -> None:
        """Tabla con encabezado azul marino, filas alternadas, encabezado repetido y filas que no se parten entre páginas.

        Las tablas cortas (12 filas o menos) se mantienen enteras en una sola página (``mantener``)."""
        self.n_tab += 1
        if pie:
            cap = self.doc.add_paragraph(style="Caption")
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(3)
            r = cap.add_run(f"Tabla {self.n_tab}. ")
            r.bold = True
            _runs_con_formato(cap, pie)
        t = self.doc.add_table(rows=1 + len(filas), cols=len(encabezados))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordes_tabla(t)
        _margen_celdas(t)
        alinear = alinear or ["izq"] + ["der"] * (len(encabezados) - 1)

        def poner(celda, texto, negrita=False, color=None, alin="izq", fondo=None):
            if fondo:
                _sombrear(celda, fondo)
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par = celda.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.05
            par.alignment = {"izq": WD_ALIGN_PARAGRAPH.LEFT, "der": WD_ALIGN_PARAGRAPH.RIGHT, "centro": WD_ALIGN_PARAGRAPH.CENTER}[alin]
            _runs_con_formato(par, str(texto), size=fuente, color=color, negrita=negrita)

        for j, h in enumerate(encabezados):
            poner(t.rows[0].cells[j], h, True, RGBColor(255, 255, 255), "izq" if alinear[j] == "izq" else "centro", HEX_NAVY)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for i, fila in enumerate(filas, start=1):
            trPr = t.rows[i]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            fondo = "E8ECF7" if resaltar_fila and (i - 1) in resaltar_fila else (HEX_ZEBRA if i % 2 == 0 else None)
            for j, v in enumerate(fila):
                poner(t.rows[i].cells[j], v, negrita=(negrita_col0 and j == 0) or (resaltar_fila is not None and (i - 1) in resaltar_fila), alin=alinear[j], fondo=fondo)
        _ancho_fijo(t, anchos_cm)
        if mantener is None:
            mantener = len(filas) <= 12
        if mantener:
            for fila in t.rows[:-1]:
                for celda in fila.cells:
                    for par in celda.paragraphs:
                        par.paragraph_format.keep_with_next = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def figura(self, ruta: Path, ancho_cm: float, pie: str, alt: str | None = None) -> None:
        self.n_fig += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(2)
        with Image.open(ruta) as im:
            ar = im.height / im.width
        r = par.add_run()
        pic = r.add_picture(str(ruta), width=Cm(ancho_cm))
        pic._inline.docPr.set("descr", alt or pie)
        pic._inline.docPr.set("title", f"Figura {self.n_fig}")
        cap = self.doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figura {self.n_fig}. ")
        r.bold = True
        _runs_con_formato(cap, pie)

    def salto(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def indice(self) -> None:
        self.titulo_preliminar("Índice")
        par = self.doc.add_paragraph()
        _campo(par, 'TOC \\o "1-2" \\h \\z \\u', "Actualice el campo (clic derecho > Actualizar campo) para ver el índice.")

    def guardar(self, ruta: Path) -> None:
        ruta.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(ruta)
