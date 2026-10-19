"""Ayudantes para construir el sitio web del proyecto: texto con mini-marcado, componentes HTML y recorte de figuras."""
from __future__ import annotations

import html
import re
from pathlib import Path

from PIL import Image


def e(t) -> str:
    """Escapa texto para HTML."""
    return html.escape(str(t), quote=True)


class Raw(str):
    """Texto que ya es HTML (no se vuelve a escapar)."""


_INL = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def md(t) -> str:
    """Mini-marcado a HTML: **negrita**, *cursiva* y `código`; todo lo demás se escapa."""
    if isinstance(t, Raw):
        return str(t)
    out = []
    for p in _INL.split(str(t)):
        if not p:
            continue
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            out.append(f"<strong>{md(p[2:-2])}</strong>")
        elif p.startswith("`") and p.endswith("`") and len(p) > 2:
            out.append(f"<code>{e(p[1:-1])}</code>")
        elif p.startswith("*") and p.endswith("*") and len(p) > 2:
            out.append(f"<em>{e(p[1:-1])}</em>")
        else:
            out.append(e(p))
    return "".join(out)


# --------------------------------------------------------------------------- #
# figuras
# --------------------------------------------------------------------------- #
FUENTE_FIG = "Fuente: INEGI, ENOE 2025-T4 (SDEMT y COE2T); estimaciones ponderadas con fac_tri; IC 95 % por bootstrap de Rao-Wu."


def preparar_figura(origen: Path, destino: Path, arriba_in: float = 0.0, abajo_in: float = 0.0, dpi: int = 200) -> None:
    """Copia una figura recortando el encabezado (título y subtítulo, que el sitio ya dice en texto) y el pie."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(origen) as im:
        caja = (0, int(arriba_in * dpi), im.width, im.height - int(abajo_in * dpi))
        im.crop(caja).save(destino, optimize=True)


def fig(nombre: str, alt: str, pie: str, porque: str | None = None, fuente: bool = True) -> str:
    """Figura con pie y, opcionalmente, la justificación del tipo de gráfico (lo que pide la guía: gráficos justificados y honestos)."""
    cuerpo = (f'<figure class="fig"><img src="img/{e(nombre)}" alt="{e(alt)}" loading="lazy" decoding="async">'
              f'<figcaption>{md(pie)}{" " + e(FUENTE_FIG) if fuente else ""}</figcaption></figure>')
    if porque:
        cuerpo += nota("Por qué este gráfico", porque)
    return cuerpo


def nota(titulo: str, texto, tipo: str = "") -> str:
    """Bloque con etiqueta que dice qué clase de bloque es. texto puede ser un str o una lista de párrafos."""
    parrafos = texto if isinstance(texto, (list, tuple)) else [texto]
    clase = {"sencillo": " nota--sencillo", "cuidado": " nota--cuidado", "ok": " nota--ok"}.get(tipo, "")
    return f'<aside class="nota{clase}"><span class="nota__t">{e(titulo)}</span>' + "".join(f"<p>{md(p)}</p>" for p in parrafos) + "</aside>"


def cifras(items) -> str:
    """Fila de números grandes: [(valor, etiqueta), ...]."""
    return '<ul class="cifras">' + "".join(f'<li><b{" class=larga" if len(v) > 12 else ""}>{e(v)}</b><span>{md(l)}</span></li>' for v, l in items) + "</ul>"


def tabla(encabezados, filas, num=(), centro=(), pie: str | None = None, destacar=()) -> str:
    """Tabla dentro de un contenedor con desplazamiento. num/centro: índices de columna; destacar: índices de fila."""
    th = "".join(f"<th{' class=\"n\"' if j in num else ' class=\"c\"' if j in centro else ''}>{md(h)}</th>" for j, h in enumerate(encabezados))
    trs = []
    for i, f in enumerate(filas):
        tds = []
        for j, c in enumerate(f):
            cl = " class=\"n\"" if j in num else (" class=\"c\"" if j in centro else "")
            tds.append(f"<td{cl}>{md(c)}</td>")
        trs.append(f"<tr{' class=\"res\"' if i in destacar else ''}>{''.join(tds)}</tr>")
    cap = f"<caption>{md(pie)}</caption>" if pie else ""
    return f'<div class="scroll"><table class="tabla">{cap}<thead><tr>{th}</tr></thead><tbody>{"".join(trs)}</tbody></table></div>'


def si_no(valor: bool) -> Raw:
    return Raw('<span class="si">Sí</span>' if valor else '<span class="no">No</span>')


def codigo(texto: str, lenguaje: str = "") -> str:
    return f'<pre class="codigo" data-lenguaje="{e(lenguaje)}"><code>{e(texto)}</code></pre>'


def pasos(items) -> str:
    return '<ol class="pasos">' + "".join(f"<li><div>{md(i)}</div></li>" for i in items) + "</ol>"


def tarjeta(titulo: str, cuerpo_html: str) -> str:
    return f'<article class="tarjeta"><h3>{md(titulo)}</h3>{cuerpo_html}</article>'


def lista(items) -> str:
    return "<ul>" + "".join(f"<li>{md(i)}</li>" for i in items) + "</ul>"


def p(texto) -> str:
    return f"<p>{md(texto)}</p>"


def dos(izq: str, der: str, clase: str = "") -> str:
    return f'<div class="dos {clase}"><div>{izq}</div><div>{der}</div></div>'


def barras(items, div: bool = False, escala: float = 100.0) -> str:
    """Barras horizontales hechas con HTML: [(etiqueta, valor, texto_valor, extra)]. div=True → barras que parten del centro (positivas y negativas)."""
    filas = []
    for it in items:
        etiqueta, valor, texto = it[0], float(it[1]), it[2]
        extra = it[3] if len(it) > 3 else ""
        pct = min(100.0, abs(valor) / escala * 100.0)
        clase = " neg" if (div and valor < 0) else (" " + extra if extra else "")
        filas.append(f'<div class="barra{" barra--div" if div else ""}"><span>{md(etiqueta)}</span><span class="pista"><i class="{clase.strip()}" style="--v:{pct:.2f}"></i></span><b>{e(texto)}</b></div>')
    return '<div class="barras" role="img" aria-label="Gráfica de barras con los valores en el texto">' + "".join(filas) + "</div>"


# --------------------------------------------------------------------------- #
# guiones (tomados de las notas del orador de la presentación completa) y escenas
# --------------------------------------------------------------------------- #
class Guiones:
    """Lee el título, las notas del orador, quién presenta y la marca ★ de cada diapositiva de la presentación completa."""

    def __init__(self, ruta: Path):
        self.items: list[dict] = []
        if not Path(ruta).exists():
            return
        from pptx import Presentation
        for s in Presentation(str(ruta)).slides:
            titulo = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
            notas = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            quien = None
            m = re.match(r"\[(.+?)\]\n", notas)
            if m:
                q = re.search(r"presenta (.+)$", m.group(1))
                quien = q.group(1) if q else None
                notas = notas[m.end():]
            estrella = any(sh.name == "Marca ruta corta" for sh in s.shapes)
            self.items.append(dict(titulo=titulo, notas=notas.strip(), quien=quien, estrella=estrella))

    @staticmethod
    def _n(t: str) -> str:
        return t.replace("–", "-")

    def de(self, prefijos):
        """prefijos: comienzos de título de las diapositivas que respaldan la escena → (párrafos de guion, quien presenta, ¿ruta corta?)."""
        textos, quien, corta = [], None, False
        for pref in prefijos:
            it = next((x for x in self.items if self._n(x["titulo"]).startswith(self._n(pref))), None)
            if it is None:
                if not self.items:
                    continue
                raise KeyError(f"La presentación completa no tiene una diapositiva que empiece con: {pref!r}")
            textos.append(it["notas"])
            quien = quien or it["quien"]
            corta = corta or it["estrella"]
        return textos, quien, corta


class Ctx:
    """Contexto de construcción: datos, guiones, capítulos y registro de escenas (para el índice y el modo presentación)."""

    CAPS = [("Inicio", "", ""), ("Selección", "KDD 1", "Sem. 1"), ("Preprocesamiento", "KDD 2", "Sem. 4"),
            ("Transformación e ingeniería de datos", "KDD 3", "Sem. 4 · 6 · 7"), ("Minería y análisis", "KDD 4", "Sem. 2 · 3 · 5"), ("Interpretación", "KDD 5", "Sem. 5")]

    def __init__(self, D, G: Guiones):
        self.D, self.G = D, G
        self.reg: list[dict] = []

    def etiquetas(self, cap: int, sub: str = "", corta: bool = False) -> str:
        nombre, fase, sem = self.CAPS[cap]
        et = f"<span>{e(nombre)}</span>"
        if fase:
            et += f'<span class="sem">{e(fase)} · {e(sem)}</span>'
        if sub:
            et += f"<span>{e(sub)}</span>"
        if corta:
            et += '<span class="corta">★ ruta corta</span>'
        return et

    @staticmethod
    def guion_html(textos, quien) -> str:
        if not textos:
            return ""
        cab = f"Guion para exponer{' · propuesta: presenta ' + e(quien) if quien else ''}"
        return f'<details class="guion"><summary>{cab}</summary>' + "".join(f'<p class="guion__txt">{e(t)}</p>' for t in textos) + "</details>"

    def atributos(self, id: str, cap: int, titular: str, quien, corta: bool) -> str:
        self.reg.append(dict(id=id, cap=cap, titulo=titular, corta=corta))
        atr = f'id="{id}" data-escena data-cap="{e(self.CAPS[cap][0])}" data-titulo="{e(titular)}"'
        atr += f' data-quien="{e(quien)}"' if quien else ""
        atr += ' data-corta="1"' if corta else ""
        return atr

    def escena(self, id: str, cap: int, titular: str, cuerpo: str, g=(), guion: str | None = None, sub: str = "", lead: str = "", clase: str = "") -> str:
        textos, quien, corta = self.G.de(g) if g else ([], None, False)
        if guion:
            textos = [guion] + textos
        leadh = f'<p class="lead">{md(lead)}</p>' if lead else ""
        return (f'<section class="escena {clase}" {self.atributos(id, cap, titular, quien, corta)}><div class="cont">'
                f'<p class="etq">{self.etiquetas(cap, sub, corta)}</p><h2 class="titular">{md(titular)}</h2>{leadh}{cuerpo}{self.guion_html(textos, quien)}</div></section>')


# --------------------------------------------------------------------------- #
# color de los mosaicos (rampa secuencial de un solo tono: la misma lógica de las gráficas del proyecto)
# --------------------------------------------------------------------------- #
_RAMPA = ["#d3e3fb", "#9dc0f2", "#5b97e6", "#1f5fb8", "#0b2c66"]


def _hex(c: str) -> tuple[int, int, int]:
    return int(c[1:3], 16), int(c[3:5], 16), int(c[5:7], 16)


def color_rampa(t: float) -> tuple[str, str]:
    """t en [0, 1] → (color de fondo, color de texto legible)."""
    t = max(0.0, min(1.0, t))
    x = t * (len(_RAMPA) - 1)
    i = min(int(x), len(_RAMPA) - 2)
    f = x - i
    a, b = _hex(_RAMPA[i]), _hex(_RAMPA[i + 1])
    rgb = tuple(round(a[k] + (b[k] - a[k]) * f) for k in range(3))
    lum = (0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]) / 255
    return "#%02x%02x%02x" % rgb, ("#0b1230" if lum > 0.55 else "#ffffff")
