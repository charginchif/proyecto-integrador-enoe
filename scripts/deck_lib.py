"""Biblioteca compartida de las presentaciones (corta y completa): tema, cifras, ayudantes de diseño y la clase Deck.

Todas las cifras se leen de results/ (resumen.json y CSV); las figuras son las de figures/.
Tema propio (colores y fuentes en el tema del archivo), gráficas nativas editables, texto alternativo en imágenes y notas del orador.
"""
from __future__ import annotations

import json
import sys
import uuid
from pathlib import Path

import pandas as pd
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR as TC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from enoe import config as C  # noqa: E402

FIG = C.FIGURES
RES = C.RESULTS
R = json.loads((RES / "resumen.json").read_text(encoding="utf-8"))


def leer(n):
    return pd.read_csv(RES / f"{n}.csv", encoding="utf-8-sig")


# ------------------------------- cifras ------------------------------- #
bn = leer("brecha_nivel").set_index("nivel")
tot = bn.loc["Total"]
esc = leer("escalera_brecha").set_index("modelo")
cu = leer("brecha_cuantiles").set_index("percentil")
cal = leer("calidad_ingreso")
prem = leer("premio_por_nivel").set_index("nivel")
ext = leer("etl_validacion_externa")
deff = leer("efecto_diseno").set_index("estimador")
robr = leer("robustez_retorno")
rob = leer("robustez_brecha")
til = leer("til1_entidad")
ret = leer("retorno_entidad")
bit = leer("etl_bitacora").set_index("id")
rt, til1, ind = R["retorno"], R["til1"], R["individual_informalidad"]
A_t, B_t, C_t, D_t = (float(cal[(cal.grupo == g) & (cal.poblacion == "Total")]["pct"].iloc[0]) for g in "ABCD")
A_n1 = float(cal[(cal.grupo == "A") & (cal.poblacion == "Nivel 1. Primaria incompleta")]["pct"].iloc[0])
A_n4 = float(cal[(cal.grupo == "A") & (cal.poblacion == "Nivel 4. Medio superior y superior")]["pct"].iloc[0])
g5 = esc.loc["G5 + rama de actividad"]
dq = rt["dif_q4_q1_log"]
rho_max, rho_min, ret_max, ret_min = rt["rho_max"], rt["rho_min"], rt["ret_max"], rt["ret_min"]
mx, mn = til.loc[til.til1_est.idxmax()], til.loc[til.til1_est.idxmin()]
rob_tot = rob[rob.grupo == "Total"].set_index("especificacion")
hora_specs = rob_tot[~rob_tot.index.str.startswith(("S6", "S8", "S9"))]
n_u, n_s = R["muestras"]["universo"], R["muestras"]["salarial"]

# ------------------------------- tema ------------------------------- #
THEME = dict(
    name="Proyecto ENOE", head="Calibri", body="Calibri",
    colors=dict(dk1="1B1B1F", lt1="FFFFFF", dk2="1E2A6B", lt2="F3F5FB", accent1="2A78D6", accent2="EB6834", accent3="1BAF7A",
                accent4="EDA100", accent5="6250D6", accent6="E34948", hlink="2A78D6", folHlink="6250D6"))
RGB = dict(muted="5F6368", grid="E1E0D9", soft="CADCFC", navy2="2B3A85", ink2="3C3F48")
MONO = "Courier New"


def rgb(h):
    return RGBColor.from_string(h)


NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
M = 0.6                                   # margen (pulgadas)
SW, SH = 13.333, 7.5


def aplicar_tema(prs):
    parte = prs.slide_master.part.part_related_by(RT.THEME)
    raiz = etree.fromstring(parte.blob)
    a = "{%s}" % NS["a"]
    esq = raiz.find(".//a:clrScheme", NS)
    esq.set("name", THEME["name"])
    for k, v in THEME["colors"].items():
        nodo = esq.find(f"a:{k}", NS)
        for h in list(nodo):
            nodo.remove(h)
        etree.SubElement(nodo, a + "srgbClr").set("val", v)
    fs = raiz.find(".//a:fontScheme", NS)
    fs.set("name", THEME["name"])
    fs.find("a:majorFont/a:latin", NS).set("typeface", THEME["head"])
    fs.find("a:minorFont/a:latin", NS).set("typeface", THEME["body"])
    raiz.set("name", THEME["name"])
    parte._blob = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)


# ------------------------------- ayudantes ------------------------------- #
def color(fmt, spec):
    """spec: enum de color de tema, o texto hex RGB."""
    if isinstance(spec, str):
        fmt.rgb = rgb(spec)
    else:
        fmt.theme_color = spec


def estilo_run(run, size, bold=False, spec=TC.TEXT_1, italic=False, fuente=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if fuente:
        run.font.name = fuente
    color(run.font.color, spec)


def caja(slide, x, y, w, h, nombre, ancla=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = nombre
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancla
    return tb, tf


def parrafos(tf, items, size=16, spec=TC.TEXT_1, bold=False, espacio=6, alinear=PP_ALIGN.LEFT, vinetas=False, fuente=None):
    """items: lista de str o de listas de (texto, {bold, spec, size}) para runs mezclados."""
    primero = True
    for it in items:
        p = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        p.alignment = alinear
        p.space_after = Pt(espacio)
        runs = [(it, {})] if isinstance(it, str) else it
        for texto, op in runs:
            r = p.add_run()
            r.text = texto
            estilo_run(r, op.get("size", size), op.get("bold", bold), op.get("spec", spec), op.get("italic", False), op.get("fuente", fuente))
        if vinetas:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.28)))
            pPr.set("indent", str(-Inches(0.28)))
            for hijo in pPr.findall("{%s}buNone" % NS["a"]):
                pPr.remove(hijo)
            bu = etree.SubElement(pPr, "{%s}buChar" % NS["a"])
            bu.set("char", "•")


def titulo(slide, texto, oscuro=False):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(M), Inches(0.38), Inches(SW - 2 * M), Inches(1.1)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = texto
    estilo_run(r, 32, True, TC.BACKGROUND_1 if oscuro else TC.TEXT_2)


def tarjeta(slide, x, y, w, h, nombre, fondo=TC.BACKGROUND_2, forma=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.06):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = nombre
    s.shadow.inherit = False
    s.line.fill.background()
    s.fill.solid()
    color(s.fill.fore_color, fondo)
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    s.text_frame.text = ""
    return s


def texto_en_forma(shape, items, size=16, spec=TC.TEXT_1, bold=False, alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE, margen=0.12, espacio=2, fuente=None):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    tf.margin_left = tf.margin_right = Inches(margen)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    parrafos(tf, items, size, spec, bold, espacio, alinear, fuente=fuente)


def estadistica(slide, x, y, w, h, valor, etiqueta, nombre, vsize=40, lsize=14, fondo=TC.BACKGROUND_2, vspec=TC.TEXT_2, lspec=TC.TEXT_1):
    tarjeta(slide, x, y, w, h, nombre, fondo)
    _, tf = caja(slide, x + 0.2, y + 0.1, w - 0.4, h - 0.2, nombre + " texto", MSO_ANCHOR.MIDDLE)
    parrafos(tf, [[(valor, {"size": vsize, "bold": True, "spec": vspec})]], espacio=0)
    parrafos_extra = tf.add_paragraph()
    r = parrafos_extra.add_run()
    r.text = etiqueta
    estilo_run(r, lsize, False, lspec)


CORTES = ROOT / "presentacion" / "img" / "cortes"


def recortar(ruta, arriba_in=0.0, abajo_in=0.0, dpi=200, caja_px=None):
    """Recorta el encabezado (título/subtítulo) y el pie de una figura para que no repita el título de la diapositiva."""
    CORTES.mkdir(parents=True, exist_ok=True)
    destino = CORTES / Path(ruta).name
    with Image.open(ruta) as im:
        if caja_px is None:
            caja_px = (0, int(arriba_in * dpi), im.width, im.height - int(abajo_in * dpi))
        im.crop(caja_px).save(destino)
    return destino


def imagen(slide, nombre_png, x, y, w, h, alt, ruta=None, cortar=None):
    ruta = Path(ruta) if ruta else FIG / nombre_png
    if cortar is not None:
        ruta = recortar(ruta, *cortar) if not isinstance(cortar, dict) else recortar(ruta, **cortar)
    with Image.open(ruta) as im:
        iw, ih = im.size
    escala = min(w / iw, h / ih)
    pw, ph = iw * escala, ih * escala
    pic = slide.shapes.add_picture(str(ruta), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    pic.name = nombre_png
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def pie(slide, n, oscuro=False):
    spec = RGB["soft"] if oscuro else RGB["muted"]
    _, tf = caja(slide, M, SH - 0.42, 9.5, 0.28, "Fuente", MSO_ANCHOR.MIDDLE)
    parrafos(tf, [[("Fuente: INEGI, ENOE 2025-T4 (SDEMT y COE2T) · cálculos del equipo", {"size": 10, "spec": spec})]], espacio=0)
    _, tf = caja(slide, SW - M - 1.0, SH - 0.42, 1.0, 0.28, "Número de diapositiva", MSO_ANCHOR.MIDDLE)
    parrafos(tf, [[(str(n), {"size": 10, "spec": spec})]], espacio=0, alinear=PP_ALIGN.RIGHT)


def fondo(slide, oscuro):
    slide.background.fill.solid()
    color(slide.background.fill.fore_color, TC.TEXT_2 if oscuro else TC.BACKGROUND_1)


def punto_numerado(slide, x, y, d, num, nombre):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.name = nombre
    c.shadow.inherit = False
    c.line.fill.background()
    c.fill.solid()
    color(c.fill.fore_color, TC.ACCENT_1)
    texto_en_forma(c, [str(num)], 18, TC.BACKGROUND_1, True, margen=0)


def estilo_grafica(chart, size=14):
    chart.font.size = Pt(size)
    color(chart.font.color, TC.TEXT_1)


def eje_valores(chart, formato):
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(RGB["grid"])
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(14)
    va.tick_labels.font.color.rgb = rgb(RGB["muted"])
    va.tick_labels.number_format = formato
    va.tick_labels.number_format_is_linked = False
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb("C3C2B7")
    ca.tick_labels.font.size = Pt(14)
    ca.tick_labels.font.color.rgb = rgb(RGB["ink2"])
    ca.has_major_gridlines = False


# ------------------------------- ayudantes nuevos (versión completa) ------------------------------- #
def tabla_nativa(slide, x, y, anchos, encabezados, filas, nombre, size=13, alto_fila=0.42, alinear=None, negrita_col0=False, resaltar=(), alto_enc=None):
    """Tabla nativa con encabezado azul marino y filas alternadas. anchos en pulgadas; filas: lista de listas de texto."""
    alinear = alinear or [PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * (len(encabezados) - 1)
    alto_enc = alto_enc or alto_fila
    gf = slide.shapes.add_table(len(filas) + 1, len(encabezados), Inches(x), Inches(y), Inches(sum(anchos)), Inches(alto_enc + alto_fila * len(filas)))
    gf.name = nombre
    t = gf.table
    for j, (h, w) in enumerate(zip(encabezados, anchos)):
        t.columns[j].width = Inches(w)
        c = t.cell(0, j)
        c.fill.solid()
        color(c.fill.fore_color, TC.TEXT_2)
        c.text_frame.text = ""
        r = c.text_frame.paragraphs[0].add_run()
        r.text = h
        estilo_run(r, size, True, TC.BACKGROUND_1)
        c.text_frame.paragraphs[0].alignment = alinear[j]
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    t.rows[0].height = Inches(alto_enc)
    for i, fila in enumerate(filas, start=1):
        t.rows[i].height = Inches(alto_fila)
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.fill.solid()
            color(c.fill.fore_color, TC.BACKGROUND_2 if i % 2 == 0 else TC.BACKGROUND_1)
            c.text_frame.text = ""
            c.text_frame.word_wrap = True
            r = c.text_frame.paragraphs[0].add_run()
            r.text = str(val)
            estilo_run(r, size, (negrita_col0 and j == 0) or (i - 1) in resaltar, TC.TEXT_1)
            c.text_frame.paragraphs[0].alignment = alinear[j]
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.08)
    return t


def panel_codigo(slide, x, y, w, h, nombre, lineas, size=12, titulo_txt=None):
    """Panel de código con fondo claro y fuente monoespaciada (Courier New)."""
    tarjeta(slide, x, y, w, h, nombre, fondo="F2F3F7", radio=0.03)
    if titulo_txt:
        _, tf = caja(slide, x + 0.2, y + 0.08, w - 0.4, 0.3, nombre + " título", MSO_ANCHOR.MIDDLE)
        parrafos(tf, [[(titulo_txt, {"size": 13, "bold": True, "spec": TC.TEXT_2})]], espacio=0)
        y += 0.4
        h -= 0.4
    _, tf = caja(slide, x + 0.2, y + 0.1, w - 0.4, h - 0.2, nombre + " texto", MSO_ANCHOR.TOP)
    parrafos(tf, [[(ln if ln else " ", {"size": size, "fuente": MONO, "spec": "2B2F3A"})] for ln in lineas], espacio=0)


def etiqueta(slide, texto, x, y, w, h, nombre, fondo=TC.ACCENT_1, spec=TC.BACKGROUND_1, size=11, negrita=True):
    sh = tarjeta(slide, x, y, w, h, nombre, fondo=fondo, radio=0.5)
    texto_en_forma(sh, [[(texto, {"size": size, "bold": negrita, "spec": spec})]], margen=0.04, espacio=0)
    return sh


def activar_secciones(prs, bloques):
    """Agrega secciones de PowerPoint (panel de miniaturas): bloques = [(nombre, [índices base 0 de diapositivas]), ...]."""
    ids = [s.get("id") for s in prs.slides._sldIdLst]
    p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    pres = prs.part._element
    ext_lst = pres.find("{%s}extLst" % pns)
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, "{%s}extLst" % pns)
    ext = etree.SubElement(ext_lst, "{%s}ext" % pns)
    ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    lst = etree.SubElement(ext, "{%s}sectionLst" % p14, nsmap={"p14": p14})
    for nombre, idx in bloques:
        sec = etree.SubElement(lst, "{%s}section" % p14)
        sec.set("name", nombre)
        sec.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        sl = etree.SubElement(sec, "{%s}sldIdLst" % p14)
        for i in idx:
            etree.SubElement(sl, "{%s}sldId" % p14).set("id", ids[i])


class Deck:
    """Presentación en construcción. La numeración, el pie, el rótulo de bloque y la marca «ruta corta» se manejan aquí."""

    def __init__(self, salida: Path, completa: bool = False):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        aplicar_tema(self.prs)
        self.prs.core_properties.title = "Radiografía socioeconómica de México: escolaridad, ingreso y brecha de género"
        self.prs.core_properties.author = "Equipo 1: Juan Evelio Culin Grajeda, Angel Yireh Garcia Jara, Luis Andre Gutierrez Zuñiga"
        self.prs.core_properties.subject = "Proyecto Integrador de Ciencia de Datos · ENOE 2025-T4 (INEGI)" + (" · presentación completa" if completa else "")
        self.layout = self.prs.slide_layouts[5]      # «Solo el título»
        self.salida = Path(salida)
        self.completa = completa
        self.bloque_actual: str | None = None
        self.presenta: str | None = None
        self.bloques: list[tuple[str, list[int]]] = []

    @property
    def n(self) -> int:
        return len(self.prs.slides)

    def bloque(self, nombre: str, presenta: str | None = None) -> None:
        """Las diapositivas siguientes pertenecen a este bloque (sección, rótulo superior y quién presenta)."""
        self.bloque_actual = nombre
        self.presenta = presenta
        self.bloques.append((nombre, []))

    def nueva(self, titulo_txt, notas, oscuro=False, con_pie=True):
        s = self.prs.slides.add_slide(self.layout)
        fondo(s, oscuro)
        titulo(s, titulo_txt, oscuro)
        if con_pie:
            pie(s, self.n, oscuro)
        prefijo = f"[{self.bloque_actual}" + (f" · presenta {self.presenta}" if self.presenta else "") + "]\n" if (self.completa and self.bloque_actual) else ""
        s.notes_slide.notes_text_frame.text = prefijo + notas
        if self.bloques:
            self.bloques[-1][1].append(self.n - 1)
        if self.completa and self.bloque_actual and con_pie:
            rotulo = self.bloque_actual.upper() + (f"  ·  {self.presenta.upper()}" if self.presenta else "")
            _, tf = caja(s, M, 0.12, 8.5, 0.26, "Bloque", MSO_ANCHOR.MIDDLE)
            parrafos(tf, [[(rotulo, {"size": 11, "bold": True, "spec": RGB["soft"] if oscuro else TC.ACCENT_1})]], espacio=0)
        return s

    def ruta_corta(self) -> None:
        """Marca la última diapositiva creada como parte de la ruta corta (★)."""
        s = self.prs.slides[-1]
        if self.completa:
            etiqueta(s, "★ ruta corta", SW - M - 1.35, 0.1, 1.35, 0.3, "Marca ruta corta", fondo=TC.ACCENT_4, spec=TC.TEXT_1, size=11)

    def guardar(self) -> Path:
        if self.completa:
            activar_secciones(self.prs, [(nombre, idx) for nombre, idx in self.bloques if idx])
        self.salida.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(self.salida)
        print(f"Presentación guardada: {self.salida} ({len(self.prs.slides)} diapositivas, {self.salida.stat().st_size / 1e6:.1f} MB)")
        return self.salida
